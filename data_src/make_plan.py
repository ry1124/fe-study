# -*- coding: utf-8 -*-
"""基本情報技術者試験 1週間合格プランを作る。
出力: data_src/plan.json(アプリ用) と 基本情報\\基本情報_1週間の計画.xlsx(人が見る用)。
方針: 1日8時間。1〜4日目でテキストを1周しながら科目Bのトレースを毎日解く / 3日目と5日目に関門 / 5・6日目に模擬試験 / 7日目に受験(CBT)。
日付は持たない(1日目をいつにするかはアプリで選ぶ)。"""
import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT_JSON = HERE / "plan.json"
OUT_XLSX = HERE.parent.parent / "基本情報" / "基本情報_1週間の計画.xlsx"
DRILL = "科目Aドリル20問（アプリの「今日」から。30分）"


def g(gid, label, typ, target, **kw):
    # typ: rate=問題の正答率(自動。chs=科目Aの章 / bids=科目Bの問題) / mock=模擬試験の直近の点数(自動) / score=本番の点数(入力)
    return dict({"id": gid, "label": label, "type": typ, "target": target}, **kw)


def b(*nums, s=()):
    return [f"b{n:02d}" for n in nums] + [f"s{n:02d}" for n in s]


plan = [
    {"n": 1, "phase": "インプット", "theme": "予約・基礎理論・アルゴリズム・擬似言語（第1〜3章）", "chs": [1, 2, 3], "b": b(*range(1, 7)),
     "todo": ["試験を7日目で予約する（IPAの公式サイトからCBTの申込み。会場の空きと締切りを確認。20分）",
              "第1・2章を読み、1章読むごとにその章の科目A問題を解く（3時間）",
              "第3章を読み、科目B 演習b01〜b06を1問10分で解く。必ずメモにトレース表を書く（3時間）",
              "×の問題を、理由を説明できるまで解き直す（1時間）",
              DRILL],
     "memo": "⚠️ 予約は今日中に。CBT試験は2026年12月28日以降に休止の予定なので、それまでに受ける"},
    {"n": 2, "phase": "インプット", "theme": "コンピュータ・システム構成・データベース（第4〜6章）", "chs": [4, 5, 6], "b": b(*range(7, 13)),
     "todo": ["昨日の「間違えた問題」を解き直す（30分）",
              "第4〜6章を読み、1章読むごとにその章の問題を解く（4時間）",
              "科目B 演習b07〜b12を1問10分で解き、解説のトレースを自分で再現する（2.5時間）",
              DRILL,
              "稼働率・キャッシュ・SQLの計算問題を、答えを見ずに解き直す（30分）"],
     "memo": "第3章の擬似言語の書き方（mod・要素番号は1から・and/orの優先順位）は、毎朝5分見直す"},
    {"n": 3, "phase": "インプット", "theme": "ネットワーク・セキュリティ・科目Bセキュリティ（第7・8・12章）", "chs": [7, 8, 12], "b": b(*range(13, 19), s=range(1, 5)),
     "todo": [DRILL,
              "第7・8章を読み、1章読むごとにその章の問題を解く（3.5時間）",
              "第12章を読み、科目B セキュリティ s01〜s04を1問5分で解く（1時間）",
              "科目B 演習b13〜b18を1問8分で解く（2時間）",
              "「間違えた問題」を解き直す（1時間）"],
     "memo": "届かなければ4日目の朝1時間を、第1〜8章の×の解き直しにあてる",
     "gate": [g("g1-a", "科目A 第1〜8章の正答率", "rate", 70, chs=[1, 2, 4, 5, 6, 7, 8]),
              g("g1-b", "科目B b01〜b18・s01〜s04の正答率", "rate", 60, bids=b(*range(1, 19), s=range(1, 5)))]},
    {"n": 4, "phase": "インプット", "theme": "開発・マネジメント・ストラテジ（第9〜11章）", "chs": [9, 10, 11], "b": b(*range(19, 25), s=range(5, 9)),
     "todo": [DRILL,
              "第9〜11章を読み、1章読むごとにその章の問題を解く（3.5時間）",
              "科目B 演習b19〜b24を1問8分、セキュリティ s05〜s08を1問5分で解く（1.5時間）",
              "科目Bの×の問題を、トレース表を書いて解き直す（1.5時間）",
              "「間違えた問題」（科目A）を解き直す（1時間）"],
     "memo": "第10・11章は用語問題が多い。1問30秒で答えられるまで回す"},
    {"n": 5, "phase": "総仕上げ", "theme": "総復習・IPA公開問題・模擬試験 第1回", "chs": [], "b": [],
     "todo": [DRILL,
              "「間違えた問題」を0問にする（1.5時間）",
              "IPAの公式サイトで公開されている科目Bのサンプル問題・公開問題を、1問5分で解く（2時間）",
              "本番形式の模擬試験 第1回（アプリ。科目A 90分 → 科目B 100分）を解き、解説を読む（3.5時間）",
              "模試で間違えた所を、テキストで確認する（30分）"],
     "memo": "関門に届かなければ、6日目の午前を×の解き直しにあて、模擬試験は午後に1回だけにする",
     "gate": [g("g2-a", "科目A 全問の正答率", "rate", 75, chs=[1, 2, 4, 5, 6, 7, 8, 9, 10, 11]),
              g("g2-b", "科目B 全問の正答率", "rate", 70, bids=b(*range(1, 25), s=range(1, 9)))]},
    {"n": 6, "phase": "模擬試験", "theme": "模擬試験 第2回・弱点の総ざらい", "chs": [13], "b": [],
     "todo": ["本番形式の模擬試験 第2回（科目A 90分 → 科目B 100分）を解き、解説を読む（3.5時間）",
              "模試で間違えた問題と「間違えた問題」を、テキストと一緒に解き直す（2時間）",
              "IPAの公開問題（科目A）を、時間を計って30問解く（1.5時間）",
              "第13章「本番の解き方」を読み、直前の確認リストを全部言えるようにする（1時間）"],
     "memo": "模試の点数は関門に自動で入る。このアプリの模試は練習済みの問題から出るので、本番より高く出る。700点に届かなければ受験日を延ばすことも考える（変更の期限は申込みサイトで確認）",
     "gate": [g("g3-a", "模擬試験 科目A（直近・1000点換算）", "mock", 700, part="a"),
              g("g3-b", "模擬試験 科目B（直近・1000点換算）", "mock", 700, part="b")]},
    {"n": 7, "phase": "受験日", "theme": "受験日（CBT）", "chs": [13], "b": [],
     "todo": ["第13章の「直前の確認リスト」を読む（20分）",
              "科目Aドリル20問を解き、×を確認する（30分）",
              "科目Bの「間違えた問題」を2〜3問トレースし直して、手を慣らす（40分）",
              "受験（科目A 90分・科目B 100分）。持ち物：顔写真付きの本人確認書類、申込みの確認メール"],
     "memo": "科目Bは最後のセキュリティ4問を先に15分で解き、残り85分をアルゴリズム16問に使う。1問8分を超えたら印を付けて先へ",
     "gate": [g("g4-a", "本番 科目A（1000点満点）", "score", 600), g("g4-b", "本番 科目B（1000点満点）", "score", 600)]},
]
for p in plan:
    p.setdefault("gate", None)
    p.setdefault("memo", "")
    p["todo"] = "\n".join(f"{'①②③④⑤'[i]}{t}" for i, t in enumerate(p["todo"]))
OUT_JSON.write_text(json.dumps(plan, ensure_ascii=False, indent=0), encoding="utf-8")


def gate_text(x):
    unit = "%" if x["type"] == "rate" else "点"
    return f"{x['label']}：{x['target']}{unit}以上"


# ===== Excel(人が見る用) =====
wb = Workbook()
ws = wb.active
ws.title = "1週間の計画"
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
base = Font(name="メイリオ", size=10)
wrap = Alignment(wrap_text=True, vertical="center")
ws["A1"] = "基本情報技術者試験 1週間合格プラン（1日8時間）"
ws["A1"].font = Font(name="メイリオ", size=16, bold=True, color="12805F")
ws["A2"], ws["C2"] = "1日目の日付", "←ここに1日目の日付を入れると、日付の列が埋まります"
ws["A2"].font = Font(name="メイリオ", bold=True)
ws["B2"].number_format = "yyyy/mm/dd"
ws["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
headers = ["日目", "日付", "時期", "テーマ", "やること", "完了", "メモ", "関門"]
widths = [6, 13, 10, 36, 72, 6, 38, 36]
for c, (h, w) in enumerate(zip(headers, widths), 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.fill, cell.font = PatternFill("solid", fgColor="12805F"), Font(bold=True, color="FFFFFF", name="メイリオ")
    cell.alignment, cell.border = Alignment(horizontal="center", vertical="center"), border
    ws.column_dimensions[cell.column_letter].width = w
for i, p in enumerate(plan, 5):
    gate = "\n".join(gate_text(x) for x in (p["gate"] or []))
    vals = [f"{p['n']}日目", f'=IF($B$2="","",$B$2+{p["n"] - 1})', p["phase"], p["theme"], p["todo"], "", p["memo"], gate]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v)
        cell.font, cell.alignment, cell.border = base, wrap, border
    ws.cell(row=i, column=2).number_format = "m/d(aaa)"
    for c in (1, 2, 6):
        ws.cell(row=i, column=c).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[i].height = 17 * (p["todo"].count("\n") + 1) + 14
dv = DataValidation(type="list", formula1='"✓"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F5:F{4 + len(plan)}")
ws.freeze_panes = "A5"
ws.cell(row=6 + len(plan), column=1, value="ふだんはアプリを使う：https://ry1124.github.io/fe-study/").font = Font(name="メイリオ", color="12805F")
os.makedirs(OUT_XLSX.parent, exist_ok=True)
wb.save(OUT_XLSX)
print(f"{len(plan)}日分 → {OUT_JSON.name} / {OUT_XLSX}")
